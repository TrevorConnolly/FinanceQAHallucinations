#!/usr/bin/env python3
"""
Generate capstone PowerPoint with charts/tables from outputs/analysis/*.csv.
Run: python3 scripts/build_capstone_presentation.py
Requires: python-pptx, pandas
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "outputs" / "presentation"
ANALYSIS = ROOT / "outputs" / "analysis"

# Theme
C_BG = RGBColor(0xF8, 0xFA, 0xFC)
C_PRIMARY = RGBColor(0x1E, 0x3A, 0x5F)
C_ACCENT = RGBColor(0x2B, 0x7C, 0xA8)
C_MUTED = RGBColor(0x4A, 0x55, 0x68)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_GREEN = RGBColor(0x27, 0x6F, 0x47)
C_ORANGE = RGBColor(0xC0, 0x56, 0x21)
# Tier A delta chart: fixed series colors (negative bars stay colored, not white)
C_SERIES_CONTEXT = RGBColor(0x1E, 0x5A, 0x8E)  # blue — always for Δ context halluc.
C_SERIES_ANSWER = RGBColor(0xB9, 0x1C, 0x1C)  # red — always for Δ answer correctness


def _blank_layout(prs: Presentation):
    for idx in (6, 5, 1):
        if idx < len(prs.slide_layouts):
            return prs.slide_layouts[idx]
    return prs.slide_layouts[0]


def _set_slide_bg(slide, color: RGBColor = C_BG) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _title_band(slide, prs: Presentation, title: str, subtitle: str | None = None) -> None:
    band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        0,
        0,
        prs.slide_width,
        Inches(1.15),
    )
    band.fill.solid()
    band.fill.fore_color.rgb = C_PRIMARY
    band.line.fill.background()
    tf = band.text_frame
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.25)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.LEFT
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(14)
        p2.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)


def _body_box(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    lines: list[str],
    size: int = 15,
    color: RGBColor = C_MUTED,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(6)


def _load_data():
    cross = pd.read_csv(ANALYSIS / "tier_a_cross_corpus.csv")
    pool = pd.read_csv(ANALYSIS / "tier_a_pooled_by_industry.csv")
    tier_b = pd.read_csv(ANALYSIS / "tier_b_pooled_by_genre.csv")
    return cross, pool, tier_b


def _chart_tier_a_delta(slide, cross: pd.DataFrame, prs: Presentation) -> None:
    """Clustered column: experiments x two series."""
    df = cross.sort_values("experiment")
    cats = [str(x).replace("_", " ") for x in df["experiment"].tolist()]
    ctx = tuple(round(float(x), 4) for x in df["context_hallucination_mean_delta"].tolist())
    ans = tuple(round(float(x), 4) for x in df["answer_correctness_mean_delta"].tolist())

    chart_data = CategoryChartData()
    chart_data.categories = cats
    chart_data.add_series("Δ context halluc. (blue = this metric; lower is better)", ctx)
    chart_data.add_series("Δ answer correctness (red = this metric; higher is better)", ans)

    x, y, cx, cy = Inches(0.55), Inches(1.35), Inches(9.0), Inches(4.85)
    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data
    )
    chart = graphic_frame.chart
    chart.has_legend = True
    if chart.legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    chart.value_axis.has_major_gridlines = True
    _apply_fixed_series_colors(chart, C_SERIES_CONTEXT, C_SERIES_ANSWER)
    try:
        plot = chart.plots[0]
        plot.has_data_labels = True
    except Exception:
        pass


def _apply_fixed_series_colors(
    chart,
    rgb_first: RGBColor,
    rgb_second: RGBColor,
) -> None:
    """Force solid fills so negative columns stay blue/red (not theme 'inverted' white)."""
    rgbs = [rgb_first, rgb_second]

    def _paint_series(ser, rgb: RGBColor) -> None:
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb
        try:
            for pt in ser.points:
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb
        except Exception:
            pass

    try:
        for idx, rgb in enumerate(rgbs):
            if idx < len(chart.series):
                _paint_series(chart.series[idx], rgb)
    except Exception:
        pass
    try:
        plot = chart.plots[0]
        for idx, rgb in enumerate(rgbs):
            if idx < len(plot.series):
                _paint_series(plot.series[idx], rgb)
    except Exception:
        pass


def _chart_genre_baseline(pool: pd.DataFrame, slide, prs: Presentation) -> None:
    """Compare finance vs fda_label on baseline means for 3 metrics."""
    b = pool[pool["experiment"] == "baseline"].copy()
    if b.empty:
        return
    fin = b[b["genre"] == "finance_10k"].iloc[0]
    fda = b[b["genre"] == "fda_label"].iloc[0]

    chart_data = CategoryChartData()
    chart_data.categories = ["Recall@5", "Context halluc.", "Answer correct."]
    chart_data.add_series(
        "Finance (10-K x2)",
        (
            float(fin["recall_at_5"]),
            float(fin["context_hallucination"]),
            float(fin["answer_correctness"]),
        ),
    )
    chart_data.add_series(
        "FDA labels (x2)",
        (
            float(fda["recall_at_5"]),
            float(fda["context_hallucination"]),
            float(fda["answer_correctness"]),
        ),
    )

    x, y, cx, cy = Inches(0.6), Inches(1.35), Inches(9.0), Inches(4.9)
    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data
    )
    chart = graphic_frame.chart
    chart.has_legend = True
    if chart.legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM


def _add_table_pool(slide, pool: pd.DataFrame, prs: Presentation) -> None:
    """Small table: baseline + k3 + k8_cot for both genres, key columns."""
    want_exp = ["baseline", "k3", "k8_cot"]
    cols_show = ["genre", "experiment", "recall_at_5", "context_hallucination", "answer_correctness"]
    sub = pool[pool["experiment"].isin(want_exp)].copy()
    sub = sub.sort_values(["genre", "experiment"])

    rows = len(sub) + 1
    cols = len(cols_show)
    left, top = Inches(0.6), Inches(1.35)
    width = Inches(9.0)
    height = Inches(0.42 * rows)
    table = slide.shapes.add_table(rows, cols, left, top, width, height).table

    for j, cname in enumerate(cols_show):
        cell = table.cell(0, j)
        cell.text = cname.replace("_", " ")
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_PRIMARY
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = C_WHITE

    for i, (_, row) in enumerate(sub.iterrows(), start=1):
        for j, cname in enumerate(cols_show):
            val = row[cname]
            if isinstance(val, float):
                s = f"{val:.3f}" if abs(val) < 10 else f"{val:.2f}"
            else:
                s = str(val)
            table.cell(i, j).text = s
            for p in table.cell(i, j).text_frame.paragraphs:
                p.font.size = Pt(11)
                p.font.color.rgb = C_MUTED


def _add_tier_b_table(slide, tier_b: pd.DataFrame, prs: Presentation) -> None:
    sub = tier_b.sort_values(["genre", "experiment"])
    cols_show = ["genre", "experiment", "context_hallucination", "answer_correctness", "faithfulness"]
    rows = len(sub) + 1
    cols = len(cols_show)
    left, top = Inches(0.55), Inches(1.32)
    width = Inches(9.1)
    height = Inches(0.38 * rows)
    table = slide.shapes.add_table(rows, cols, left, top, width, height).table
    for j, cname in enumerate(cols_show):
        cell = table.cell(0, j)
        cell.text = cname.replace("_", " ")
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_ACCENT
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = C_WHITE
    for i, (_, row) in enumerate(sub.iterrows(), start=1):
        for j, cname in enumerate(cols_show):
            val = row[cname]
            s = f"{float(val):.3f}" if isinstance(val, (float, int)) else str(val)
            table.cell(i, j).text = s
            for p in table.cell(i, j).text_frame.paragraphs:
                p.font.size = Pt(10)


def _pipeline_visual(slide, prs: Presentation) -> None:
    steps = [
        ("PDF", "LlamaParse"),
        ("Chunks", "Headers + split"),
        ("Index", "Chroma + BM25 + rerank"),
        ("RAG", "GPT-4o + cites"),
        ("Eval", "Judges + RAGAS"),
    ]
    n = len(steps)
    w = Inches(1.55)
    h = Inches(0.85)
    gap = Inches(0.18)
    y = Inches(2.05)
    total_w = n * w + (n - 1) * gap
    x0 = (prs.slide_width - total_w) / 2
    for i, (t1, t2) in enumerate(steps):
        x = x0 + i * (w + gap)
        shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
        shp.fill.solid()
        shp.fill.fore_color.rgb = C_ACCENT if i % 2 == 0 else C_PRIMARY
        shp.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
        tf = shp.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = t1
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = C_WHITE
        p.alignment = PP_ALIGN.CENTER
        p2 = tf.add_paragraph()
        p2.text = t2
        p2.font.size = Pt(10)
        p2.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
        p2.alignment = PP_ALIGN.CENTER
        if i < n - 1:
            arr = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW,
                x + w + Inches(0.02),
                y + h / 2 - Inches(0.12),
                Inches(0.14),
                Inches(0.24),
            )
            arr.fill.solid()
            arr.fill.fore_color.rgb = RGBColor(0xA0, 0xAE, 0xC0)
            arr.line.fill.background()


def main() -> None:
    cross, pool, tier_b = _load_data()

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    blank = _blank_layout(prs)

    # --- 1. Title (layout 0) ---
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    _set_slide_bg(slide)
    slide.shapes.title.text = (
        "RAG Designs for Fewer Hallucinations:\nFinance & Healthcare Documents"
    )
    sub = slide.placeholders[1]
    sub.text = (
        "[Your Name]  |  [Class]  |  Adviser: [Adviser Name]\n"
        "Capstone — Finance QA Hallucinations"
    )
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = C_PRIMARY
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(32)

    # --- 2. Agenda / story (visual columns) ---
    slide = prs.slides.add_slide(blank)
    _set_slide_bg(slide)
    _title_band(slide, prs, "Talk roadmap", "What you will see")
    boxes = [
        ("1", "Why it matters", "Trust + productivity"),
        ("2", "Method", "Pipeline + controlled experiments"),
        ("3", "Data", "Charts from 4 real PDF corpora"),
        ("4", "Takeaway", "No one-size-fits-all RAG"),
    ]
    bx_w, bx_h = 2.05, 2.35
    gap = 0.35
    start_x = (10 - (4 * bx_w + 3 * gap)) / 2
    for i, (num, t, s) in enumerate(boxes):
        x = start_x + i * (bx_w + gap)
        sh = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(1.85),
            Inches(bx_w),
            Inches(bx_h),
        )
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor(0xEB, 0xF4, 0xFF) if i % 2 == 0 else RGBColor(0xF0, 0xF4, 0xF8)
        sh.line.color.rgb = C_ACCENT
        tf = sh.text_frame
        tf.text = f"{num}. {t}\n{s}"
        for j, p in enumerate(tf.paragraphs):
            p.font.size = Pt(16 if j == 0 else 13)
            p.font.color.rgb = C_PRIMARY if j == 0 else C_MUTED
            p.font.bold = j == 0

    # --- 3. Motivation ---
    slide = prs.slides.add_slide(blank)
    _set_slide_bg(slide)
    _title_band(
        slide,
        prs,
        "Motivation & goal",
        "Industry-relevant RAG for finance & healthcare text",
    )
    _body_box(
        slide,
        0.65,
        1.32,
        5.35,
        3.35,
        [
            "Professionals are adopting AI for productivity, but in finance and healthcare answers must be precise and auditable.",
            "Hallucinations are the main blocker to trust—especially when data is sensitive or regulated.",
            "Goal of this project: empirically compare RAG system designs (retrieval, prompting, citations, chunking) and measure which configurations best reduce hallucination-style failures while preserving answer quality.",
            "I also ask whether certain designs work better with different document types (e.g., SEC filings vs FDA labels) so teams can tune RAG with evidence—not guesswork.",
        ],
        14,
    )
    call = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(6.05),
        Inches(1.32),
        Inches(3.35),
        Inches(3.35),
    )
    call.fill.solid()
    call.fill.fore_color.rgb = RGBColor(0xFF, 0xFA, 0xEC)
    call.line.color.rgb = C_ORANGE
    tf = call.text_frame
    tf.word_wrap = True
    tf.text = "Industry perspective (EY)"
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.size = Pt(13)
    tf.paragraphs[0].font.color.rgb = C_ORANGE
    for line in [
        "AI & data technology consulting (EY): client concern about hallucinations and reliability on sensitive workloads.",
        "Interviews with tech risk partners: among customer complaints, trust and hallucinations rank at the top—even when efficiency benefits are clear.",
        "This work connects that practitioner concern to a reproducible evaluation harness.",
    ]:
        p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(11)
        p.font.color.rgb = C_MUTED
        p.space_after = Pt(4)
    _body_box(
        slide,
        0.65,
        4.85,
        8.75,
        2.35,
        [
            "Who benefits: any organization deploying RAG on dense technical prose—risk, compliance, clinical ops, IR, research analysts.",
        ],
        13,
    )

    # --- 4. Background & related work ---
    slide = prs.slides.add_slide(blank)
    _set_slide_bg(slide)
    _title_band(
        slide,
        prs,
        "Background & related work",
        "Where NLP consensus stands",
    )
    _body_box(
        slide,
        0.65,
        1.3,
        8.75,
        5.7,
        [
            "Consensus: LLMs can sound authoritative while stating content not supported by evidence. Retrieval-Augmented Generation (RAG) grounds answers in retrieved passages and usually reduces—but does not remove—unsupported or conflicting outputs.",
            "Research surveys (e.g., recent arXiv surveys on LLM hallucinations; RAG mitigation reviews) organize causes: bad retrieval, context overload, parametric knowledge overriding evidence, and evaluation gaps.",
            "Industry products: many vendors offer RAG stacks, guardrails, and observability—yet defaults are often tuned for generic chat, not long regulatory PDFs or filings.",
            "Gap this project fills: controlled comparison of concrete RAG knobs on real finance vs healthcare-style documents, with multi-metric evaluation (grounding judges + automated scores)—not only a single accuracy number.",
        ],
        14,
    )

    # --- 5. Implementation ---
    slide = prs.slides.add_slide(blank)
    _set_slide_bg(slide)
    _title_band(slide, prs, "Implementation", "Pipeline, experiments, and evaluation stack")
    _pipeline_visual(slide, prs)
    _body_box(
        slide,
        0.6,
        4.95,
        8.85,
        2.35,
        [
            "Ingestion: LlamaParse converts PDFs to Markdown; LangChain splits on headings + fixed window size (preset-specific for chunk ablations).",
            "Retrieval: vector store (Chroma) + BM25 + cross-encoder rerank (BGE). Generation: GPT-4o with mandatory chunk-ID citations; optional strict post-process (replace answer with “Insufficient Information.” if cites are invalid or missing).",
            "Gold questions: synthetic Q/A from sampled chunks + critique filter. Experiments: baseline (k=5), k∈{3,8,10}, chain-of-thought prompt, strict-cite, k8+CoT; Tier B varies chunking presets with regenerated gold.",
            "Evaluation: n=10 questions per run in this study; rule metrics (recall@k, citation checks), dedicated LLM judge for context-grounded hallucinations & claim rates, RAGAS (faithfulness, relevancy, correctness).",
        ],
        12,
    )

    # --- 6. Results overview (before charts) ---
    slide = prs.slides.add_slide(blank)
    _set_slide_bg(slide)
    _title_band(
        slide,
        prs,
        "Results — what you are about to see",
        "How experiments & tiers are defined",
    )
    _body_box(
        slide,
        0.65,
        1.28,
        8.75,
        5.75,
        [
            "Data: four public PDFs — two SEC-style 10-Ks and two FDA drug labels — same code path, different corpora.",
            "Tier A (same chunking & gold): fair comparison of presets against baseline—same 10 evaluation questions and gold chunk IDs. Includes baseline, top-k variants, CoT, strict citations, k8+CoT.",
            "Tier B (chunking ablations): fine / coarse / deep-header presets each rebuild the corpus and regenerate gold — not the same questions as Tier A. Interpret Tier B as “alternative chunking strategies,” not as Δ vs Tier A baseline.",
            "Charts: cross-corpus bars show mean change vs baseline averaged over all four documents. Tables pool metrics by genre (finance vs FDA labels). Colors: blue series = Δ context-hallucination flag; red = Δ answer correctness (negative values stay colored).",
        ],
        14,
    )

    # --- 7. Tier A chart: delta vs baseline ---
    slide = prs.slides.add_slide(blank)
    _set_slide_bg(slide)
    _title_band(
        slide,
        prs,
        "Tier A — Cross-corpus change vs baseline",
        "Same gold & chunks; 4 corpora averaged (see CSV for min/max)",
    )
    _chart_tier_a_delta(slide, cross, prs)
    _body_box(
        slide,
        0.55,
        6.32,
        9.1,
        1.05,
        [
            "Blue columns = Δ context-hallucination flag (lower is better). Red = Δ answer correctness (higher is better). Colors stay blue/red even when Δ is negative.",
            "Pattern: k3 shows the largest average reduction in the hallucination flag vs baseline; k8+CoT shows the largest average gain in correctness but also increases the hallucination flag on average—a tradeoff.",
        ],
        10,
        C_MUTED,
    )

    # --- 8. Table pooled by industry ---
    slide = prs.slides.add_slide(blank)
    _set_slide_bg(slide)
    _title_band(
        slide,
        prs,
        "Tier A — Pooled means (by genre)",
        "FDA labels vs 10-Ks — baseline, k3, k8+CoT",
    )
    _add_table_pool(slide, pool, prs)
    _body_box(
        slide,
        0.55,
        5.55,
        9.0,
        1.5,
        [
            "Finance: stronger recall@5 on these filings. Labels: different error profile — compare presets, not only baseline.",
        ],
        12,
    )

    # --- 9. Genre comparison chart ---
    slide = prs.slides.add_slide(blank)
    _set_slide_bg(slide)
    _title_band(
        slide,
        prs,
        "Genre contrast (Tier A baseline)",
        "Same pipeline; different documents",
    )
    _chart_genre_baseline(pool, slide, prs)
    _body_box(
        slide,
        0.55,
        6.45,
        9.0,
        0.75,
        [
            "Context halluc. = fraction of questions with flagged unsupported content (judge). Lower is better.",
        ],
        11,
    )

    # --- 10. Tier B table ---
    slide = prs.slides.add_slide(blank)
    _set_slide_bg(slide)
    _title_band(
        slide,
        prs,
        "Tier B — Chunking presets (separate gold)",
        "Compare within this tier only",
    )
    _add_tier_b_table(slide, tier_b, prs)
    _body_box(
        slide,
        0.55,
        4.55,
        9.0,
        2.5,
        [
            "Each preset re-chunks + regenerates Q/A — not comparable to Tier A deltas.",
            "Use for chunking strategy exploration; confirm with larger n.",
        ],
        13,
    )

    # --- 11. Design takeaway / synthesis ---
    slide = prs.slides.add_slide(blank)
    _set_slide_bg(slide)
    _title_band(
        slide,
        prs,
        "Design takeaway — answering the core questions",
        "What the data suggests for practitioners",
    )
    for i, (label, col) in enumerate(
        [
            ("Retrieval (k)", C_ACCENT),
            ("Prompting (CoT)", C_PRIMARY),
            ("Strict cites", C_GREEN),
        ]
    ):
        sh = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL,
            Inches(0.85 + i * 2.85),
            Inches(1.38),
            Inches(1.75),
            Inches(1.75),
        )
        sh.fill.solid()
        sh.fill.fore_color.rgb = col
        tf = sh.text_frame
        tf.text = label
        tf.paragraphs[0].font.size = Pt(11)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = C_WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    _body_box(
        slide,
        0.65,
        3.35,
        8.85,
        3.85,
        [
            "Fundamental question: which RAG design best reduces hallucinations for professional use? On average across four corpora, k=3 (narrower retrieval) delivered the largest reduction in the context-grounded hallucination flag vs baseline—suggesting smaller top-k can reduce unsupported content when the right passages are still retrieved.",
            "Tradeoff: k8 + chain-of-thought delivered the largest average gain in answer correctness vs baseline, but also the largest average increase in the hallucination flag vs baseline. So “best for correctness” and “best for grounding” are not the same preset.",
            "Industry nuance: pooled 10-Ks showed higher recall@5 than pooled FDA labels in this setup—retrieval difficulty differs by genre. No single winner for both finance-style and label-style text; teams should pick metrics (grounding vs correctness vs abstention) first, then select a preset.",
            "Strict citation post-processing: near-zero average change in the hallucination flag in the aggregate, with a small average drop in correctness—consistent with more forced abstention when cites are imperfect.",
            "Tier B (chunking): compare chunk presets within that tier only; they use different gold questions—useful for chunking strategy, not for ranking against Tier A deltas.",
        ],
        12,
        C_MUTED,
    )

    # --- 12. Limitations ---
    slide = prs.slides.add_slide(blank)
    _set_slide_bg(slide)
    _title_band(slide, prs, "Limitations & next steps", "")
    _body_box(
        slide,
        0.75,
        1.45,
        8.5,
        5.5,
        [
            "• n=10 questions / experiment — exploratory.",
            "• Four PDFs — illustrative genres, not universal finance/healthcare.",
            "• LLM judges add variance; triangulate with rule-based metrics.",
            "• Next: larger n, CIs, more rerankers & abstention policies.",
        ],
        17,
    )

    # --- 13. Conclusion ---
    slide = prs.slides.add_slide(blank)
    _set_slide_bg(slide)
    _title_band(slide, prs, "Conclusion", "")
    _body_box(
        slide,
        0.85,
        1.55,
        8.3,
        5.0,
        [
            "RAG design choices measurably shift grounding vs correctness — differently on 10-Ks vs FDA labels.",
            "Deliverable: reproducible pipeline + data-driven comparison tables for practitioners.",
            "Thank you — questions?",
        ],
        22,
        C_PRIMARY,
    )

    path = OUT / "RAG_Hallucination_Capstone.pptx"
    prs.save(path)
    print(f"Saved: {path}")


if __name__ == "__main__":
    main()
